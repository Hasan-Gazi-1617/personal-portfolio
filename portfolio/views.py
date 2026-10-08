import os
import logging
import secrets
import time
from calendar import monthrange
from datetime import datetime
from io import BytesIO

from django.contrib import messages
from django.db.models import Count, Q
from django.http import HttpResponse
from django.shortcuts import get_object_or_404, redirect, render
from django.utils import timezone
from django.views.decorators.http import require_http_methods, require_POST

logger = logging.getLogger(__name__)

from .forms import AdminTicketForm, EngineerTicketUpdateForm
from .models import ComplaintTicket, TicketActivity


def _month_bounds(raw_month):
    try:
        selected = datetime.strptime(raw_month, "%Y-%m")
    except (TypeError, ValueError):
        selected = timezone.localtime().replace(day=1, hour=0, minute=0, second=0, microsecond=0)
    start = timezone.make_aware(selected) if timezone.is_naive(selected) else selected
    last_day = monthrange(start.year, start.month)[1]
    end = start.replace(day=last_day, hour=23, minute=59, second=59, microsecond=999999)
    return start, end


def _owner_can_manage_tki(request):
    return bool(request.session.get("owner_access"))

def home(request):
    return render(request, "home.html")


def mikrotik_generator(request):
    """Interactive client-side RouterOS configuration generator."""
    return render(request, "mikrotik_generator.html")


def algorithm_studio(request):
    """Interactive data structure and algorithm learning studio."""
    return render(request, "algorithm_studio.html")


def tki_dashboard(request):
    tickets = ComplaintTicket.objects.select_related("assigned_engineer", "created_by")
    query = request.GET.get("q", "").strip()
    engineer = request.GET.get("engineer", "").strip()
    status = request.GET.get("status", "").strip()
    support = request.GET.get("support", "").strip()
    report_period = request.GET.get("period", "daily")
    if report_period not in {"daily", "monthly"}:
        report_period = "daily"

    if query:
        tickets = tickets.filter(Q(tki_id__icontains=query) | Q(client_id__icontains=query) | Q(complaint__icontains=query))
    if engineer in dict(ComplaintTicket.SupportEngineer.choices):
        tickets = tickets.filter(field_support_engineer=engineer)
    if status:
        tickets = tickets.filter(status=status)
    if support == "yes":
        tickets = tickets.exclude(field_support_engineer="")
    elif support == "no":
        tickets = tickets.filter(field_support_engineer="")

    from django.contrib.auth import get_user_model
    engineers = get_user_model().objects.filter(is_active=True, is_staff=False).order_by("first_name", "username")
    all_tickets = ComplaintTicket.objects.all()
    now = timezone.localtime()
    report_tickets = all_tickets.filter(opened_at__date=now.date()) if report_period == "daily" else all_tickets.filter(opened_at__year=now.year, opened_at__month=now.month)
    summary = {
        "total": report_tickets.count(),
        "pending": report_tickets.filter(status=ComplaintTicket.Status.PENDING).count(),
        "solved": report_tickets.filter(status=ComplaintTicket.Status.SOLVED).count(),
        "cancelled": report_tickets.filter(status=ComplaintTicket.Status.CANCELLED).count(),
        "dependency": report_tickets.exclude(dependency="").count(),
        "attention": report_tickets.filter(attention=True).count(),
    }
    workload = engineers.annotate(
        active_count=Count("assigned_tickets", filter=Q(assigned_tickets__status=ComplaintTicket.Status.PENDING)),
        closed_count=Count("assigned_tickets", filter=Q(assigned_tickets__status=ComplaintTicket.Status.SOLVED)),
    )
    engineer_report = []
    max_pending = 1
    for value, label in ComplaintTicket.SupportEngineer.choices:
        engineer_tickets = report_tickets.filter(field_support_engineer=value)
        pending_count = engineer_tickets.filter(status=ComplaintTicket.Status.PENDING).count()
        max_pending = max(max_pending, pending_count)
        engineer_report.append({
            "value": value, "label": label,
            "total": engineer_tickets.count(),
            "pending": pending_count,
            "solved": engineer_tickets.filter(status=ComplaintTicket.Status.SOLVED).count(),
        })
    for item in engineer_report:
        item["bar_percent"] = round(item["pending"] / max_pending * 100)

    dependency_report = []
    dependency_total = max(summary["dependency"], 1)
    dependency_colors = ["#32d399", "#55b8ff", "#f6bd47", "#a78bfa", "#fb7185", "#94a3b8"]
    gradient_stops = []
    cursor = 0
    for index, (value, label) in enumerate(ComplaintTicket.Dependency.choices):
        count = report_tickets.filter(dependency=value).count()
        percent = round(count / dependency_total * 100)
        color = dependency_colors[index]
        if count:
            gradient_stops.append(f"{color} {cursor}% {cursor + percent}%")
            cursor += percent
        dependency_report.append({"value": value, "label": label, "count": count, "percent": percent, "color": color})
    dependency_gradient = "conic-gradient(" + (", ".join(gradient_stops) if gradient_stops else "#334155 0 100%") + ")"
    return render(request, "tki/dashboard.html", {
        "tickets": tickets[:150], "engineers": engineers, "workload": workload,
        "summary": summary, "statuses": ComplaintTicket.Status.choices,
        "support_engineers": ComplaintTicket.SupportEngineer.choices,
        "report_period": report_period, "engineer_report": engineer_report,
        "dependency_report": dependency_report, "dependency_gradient": dependency_gradient,
        "can_edit_tki": bool(request.session.get("owner_access")),
        "current_month": now.strftime("%Y-%m"),
    })


def tki_export_xlsx(request):
    if not _owner_can_manage_tki(request):
        messages.error(request, "Owner access is required to export TKI data.")
        return redirect("owner_login")

    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.utils import get_column_letter

    start, end = _month_bounds(request.GET.get("month"))
    tickets = ComplaintTicket.objects.filter(opened_at__range=(start, end)).order_by("opened_at", "id")
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "TKI Report"
    dark = "102033"
    green = "35D69A"
    muted = "DCE7F2"
    thin = Side(style="thin", color="CAD6E2")

    sheet.merge_cells("A1:L1")
    sheet["A1"] = f"ISP TKI Monthly Report — {start.strftime('%B %Y')}"
    sheet["A1"].font = Font(size=18, bold=True, color="FFFFFF")
    sheet["A1"].fill = PatternFill("solid", fgColor=dark)
    sheet["A1"].alignment = Alignment(horizontal="center", vertical="center")
    sheet.row_dimensions[1].height = 34

    counts = {
        "Total": tickets.count(),
        "Pending": tickets.filter(status=ComplaintTicket.Status.PENDING).count(),
        "Solved": tickets.filter(status=ComplaintTicket.Status.SOLVED).count(),
        "Cancelled": tickets.filter(status=ComplaintTicket.Status.CANCELLED).count(),
    }
    for column, (label, count) in enumerate(counts.items(), 1):
        cell = sheet.cell(3, column, f"{label}: {count}")
        cell.font = Font(bold=True, color=dark)
        cell.fill = PatternFill("solid", fgColor="E8F8F1")

    headers = ["Date", "Time", "TKI ID", "Client Code", "Status", "Dependency", "Higher-Level NOC", "Field Support Engineer", "Remarks", "Aging", "Attention", "Issue"]
    for column, label in enumerate(headers, 1):
        cell = sheet.cell(5, column, label)
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor=dark)
        cell.alignment = Alignment(horizontal="center")

    def safe_text(value):
        text = str(value or "")
        return "'" + text if text.startswith(("=", "+", "-", "@")) else text

    status_colors = {"pending": "FFF2CC", "solved": "D9EAD3", "cancelled": "F4CCCC"}
    for row_number, ticket in enumerate(tickets, 6):
        local_opened = timezone.localtime(ticket.opened_at)
        values = [
            local_opened.date(), local_opened.time().replace(microsecond=0), safe_text(ticket.tki_id),
            safe_text(ticket.client_id), ticket.get_status_display(), ticket.get_dependency_display() or "None",
            ticket.get_higher_level_noc_display() or "—", ticket.get_field_support_engineer_display() or "—",
            safe_text(ticket.remarks or ticket.complaint), ticket.aging,
            "Required" if ticket.attention else "Normal", ticket.get_category_display(),
        ]
        for column, value in enumerate(values, 1):
            cell = sheet.cell(row_number, column, value)
            cell.border = Border(bottom=thin)
            cell.alignment = Alignment(vertical="top", wrap_text=column in {9, 12})
        sheet.cell(row_number, 5).fill = PatternFill("solid", fgColor=status_colors.get(ticket.status, muted))

    widths = [13, 12, 18, 16, 13, 20, 20, 23, 42, 12, 14, 22]
    for column, width in enumerate(widths, 1):
        sheet.column_dimensions[get_column_letter(column)].width = width
    sheet.freeze_panes = "A6"
    sheet.auto_filter.ref = f"A5:L{max(sheet.max_row, 5)}"
    sheet.sheet_view.showGridLines = False
    sheet.page_setup.orientation = "landscape"
    sheet.page_setup.fitToWidth = 1

    output = BytesIO()
    workbook.save(output)
    response = HttpResponse(output.getvalue(), content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
    response["Content-Disposition"] = f'attachment; filename="ISP-TKI-{start.strftime("%Y-%m")}.xlsx"'
    return response


@require_POST
def tki_delete_month(request):
    if not _owner_can_manage_tki(request):
        messages.error(request, "Owner access is required to delete TKI data.")
        return redirect("owner_login")
    if request.POST.get("confirm", "").strip().upper() != "DELETE":
        messages.error(request, "Type DELETE to confirm the monthly deletion.")
        return redirect("tki_dashboard")
    start, end = _month_bounds(request.POST.get("month"))
    queryset = ComplaintTicket.objects.filter(opened_at__range=(start, end))
    count = queryset.count()
    queryset.delete()
    messages.success(request, f"{count} TKI record(s) from {start.strftime('%B %Y')} were deleted.")
    return redirect("tki_dashboard")


def tki_detail(request, pk):
    ticket = get_object_or_404(ComplaintTicket.objects.select_related("assigned_engineer", "created_by"), pk=pk)
    form = EngineerTicketUpdateForm(instance=ticket)
    return render(request, "tki/detail.html", {
        "ticket": ticket, "form": form,
        "can_update": bool(request.session.get("owner_access")),
        "can_edit_tki": bool(request.session.get("owner_access")),
    })


@require_POST
def tki_update(request, pk):
    if not request.session.get("owner_access"):
        messages.error(request, "Owner access is required to update TKI records.")
        return redirect("owner_login")
    ticket = get_object_or_404(ComplaintTicket, pk=pk)
    form = EngineerTicketUpdateForm(request.POST, instance=ticket)
    if form.is_valid():
        updated = form.save(commit=False)
        if updated.status == ComplaintTicket.Status.SOLVED and not updated.resolved_at:
            updated.resolved_at = timezone.now()
        elif updated.status != ComplaintTicket.Status.SOLVED:
            updated.resolved_at = None
        updated.save()
        actor = request.user if request.user.is_authenticated else None
        TicketActivity.objects.create(ticket=ticket, actor=actor, action="TKI updated", detail=f"Status: {updated.get_status_display()}")
        messages.success(request, "TKI feedback updated.")
        return redirect("tki_detail", pk=ticket.pk)
    return render(request, "tki/detail.html", {
        "ticket": ticket, "form": form, "can_update": True,
        "can_edit_tki": True,
    }, status=400)


def tki_create(request):
    if not request.session.get("owner_access"):
        messages.error(request, "Owner access is required to create TKI records.")
        return redirect("owner_login")
    form = AdminTicketForm(request.POST or None)
    if request.method == "POST" and form.is_valid():
        ticket = form.save(commit=False)
        ticket.created_by = request.user if request.user.is_authenticated else None
        if ticket.assigned_engineer_id:
            ticket.support_assigned_at = timezone.now()
            if not ticket.status:
                ticket.status = ComplaintTicket.Status.PENDING
        ticket.save()
        actor = request.user if request.user.is_authenticated else None
        TicketActivity.objects.create(ticket=ticket, actor=actor, action="TKI created")
        messages.success(request, f"{ticket.tki_id} created successfully.")
        return redirect("tki_detail", pk=ticket.pk)
    return render(request, "tki/form.html", {"form": form, "title": "Create TKI"})


def tki_admin_edit(request, pk):
    if not request.session.get("owner_access"):
        messages.error(request, "Owner access is required to edit TKI records.")
        return redirect("owner_login")
    ticket = get_object_or_404(ComplaintTicket, pk=pk)
    previous_engineer_id = ticket.assigned_engineer_id
    form = AdminTicketForm(request.POST or None, instance=ticket)
    if request.method == "POST" and form.is_valid():
        updated = form.save(commit=False)
        if updated.assigned_engineer_id != previous_engineer_id:
            updated.support_assigned_at = timezone.now() if updated.assigned_engineer_id else None
        if updated.status == ComplaintTicket.Status.SOLVED and not updated.resolved_at:
            updated.resolved_at = timezone.now()
        updated.save()
        actor = request.user if request.user.is_authenticated else None
        TicketActivity.objects.create(ticket=ticket, actor=actor, action="Full TKI update")
        messages.success(request, "TKI updated successfully.")
        return redirect("tki_detail", pk=ticket.pk)
    return render(request, "tki/form.html", {"form": form, "title": f"Edit {ticket.tki_id}", "ticket": ticket})


@require_http_methods(["GET", "POST"])
def owner_login(request):
    if request.session.get("owner_access"):
        return redirect("home")

    if request.method == "POST":
        now = int(time.time())
        attempts = request.session.get("owner_login_attempts", [])
        attempts = [stamp for stamp in attempts if now - stamp < 15 * 60]
        if len(attempts) >= 5:
            messages.error(request, "Too many attempts. Please wait 15 minutes and try again.")
            request.session["owner_login_attempts"] = attempts
            return render(request, "owner_login.html", status=429)

        submitted_key = request.POST.get("access_key", "")
        configured_key = os.environ.get("OWNER_ACCESS_KEY", "")

        if configured_key and secrets.compare_digest(submitted_key, configured_key):
            request.session.cycle_key()
            request.session["owner_access"] = True
            request.session.set_expiry(60 * 60 * 24 * 7)
            request.session.pop("owner_login_attempts", None)
            return redirect("home")

        attempts.append(now)
        request.session["owner_login_attempts"] = attempts
        logger.warning(
            "Rejected owner login attempt",
            extra={"client_ip": request.META.get("REMOTE_ADDR", "unknown"), "attempt_count": len(attempts)},
        )
        messages.error(request, "Invalid owner access key.")

    return render(request, "owner_login.html")


@require_POST
def owner_logout(request):
    request.session.pop("owner_access", None)
    request.session.cycle_key()
    return redirect("home")
